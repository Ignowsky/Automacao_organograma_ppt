import io
import os
import json
from typing import List, Dict, Any, BinaryIO, Optional
from pptx import Presentation
from pptx.slide import Slide
from src.domain.models import (
    HeadcountRecord,
    ChangeLogEntry,
    ExecutionReport,
    UpdateStatus,
    TargetType,
    OrganizationalRole,
)
from src.infrastructure.pptx_engine import PptxSurgicalEngine, classify_organizational_node, normalize_text_key
from src.infrastructure.databricks_repository import parse_reference_date_info
from src.infrastructure.structure_sync import StructureSyncService


class PresentationUpdaterService:
    """
    Orquestrador 100% Automático de Atualização Cirúrgica de Apresentações PowerPoint (.pptx):
    - Independe de qualquer arquivo JSON estático (slide_mapping.json pode ser excluído);
    - Reconhece geometricamente em tempo de execução todos os elementos dos slides via StructureSyncService;
    - Atualiza nomes renomeados, exclui fisicamente áreas/unidades extintas (HC = 0 ou ausentes na base)
      e cria automaticamente novas Superintendências, Diretorias e Gerências com borda pontilhada;
    - Registra as novas caixas clonadas no Preview 16:9 (slides_preview) e oculta 100% das caixas excluídas do Canvas.
    """

    def __init__(self, mapping_json_path: str = "config/slide_mapping.json") -> None:
        if mapping_json_path and os.path.exists(mapping_json_path) and os.path.getsize(mapping_json_path) > 0:
            try:
                with open(mapping_json_path, "r", encoding="utf-8-sig") as f:
                    self.base_config: Dict[str, Any] = json.load(f)
            except Exception:
                self.base_config = {}
        else:
            self.base_config = {}
        self.config: Dict[str, Any] = {}

    @staticmethod
    def _get_shape_left_emu(slide: Slide, shape_path: str) -> Optional[int]:
        top_id = str(shape_path).split("|")[0].strip().split(".")[0].strip()
        top_sh = PptxSurgicalEngine.find_shape_by_path(slide, top_id)
        if top_sh is not None and hasattr(top_sh, "left"):
            return int(top_sh.left)
        sh = PptxSurgicalEngine.find_shape_by_path(slide, shape_path)
        if sh is not None and hasattr(sh, "left"):
            return int(sh.left)
        return None

    def _infer_segment_and_area_by_column(
        self, slide: Slide, slide_index: int, shape_key: str
    ) -> tuple[Optional[str], Optional[str]]:
        probe_path = shape_key.split("|")[0].strip() if "|" in shape_key else shape_key
        target_left = self._get_shape_left_emu(slide, probe_path)

        candidates: List[tuple[int, str, str]] = []
        for a in self.config.get("area_subtotal_shapes", []):
            if int(a["slide_index"]) == slide_index and a.get("area") != "__EXTINCT_AREA__":
                a_left = self._get_shape_left_emu(slide, str(a["shape_path"]))
                if a_left is not None and target_left is not None:
                    candidates.append((abs(a_left - target_left), a["segmento"], a["area"]))

        for u in self.config.get("unit_boxes", []):
            if int(u["slide_index"]) == slide_index and not u.get("force_remove"):
                u_left = self._get_shape_left_emu(slide, str(u["name_path"]))
                if u_left is not None and target_left is not None:
                    candidates.append((abs(u_left - target_left), u["segmento"], u["area"]))

        if candidates:
            candidates.sort(key=lambda item: item[0])
            return candidates[0][1], candidates[0][2]

        for s in self.config.get("segment_total_shapes", []):
            if int(s["slide_index"]) == slide_index:
                return s["segmento"], None

        return None, None

    def execute_update(
        self,
        pptx_stream: BinaryIO,
        records: List[HeadcountRecord],
        reference_label: str = "",
        original_filename: str = "apresentacao_original.pptx",
        custom_overrides: Optional[List[Dict[str, Any]]] = None,
        auto_layout: bool = True,
    ) -> tuple[bytes, ExecutionReport]:
        prs = Presentation(pptx_stream)
        slides_before = len(prs.slides)

        if not reference_label or not reference_label.strip():
            sample_ref = records[0].ano_mes_referencia if records else None
            reference_label = parse_reference_date_info(sample_ref)["label"]

        total_company = sum(r.headcount for r in records if r.headcount > 0)
        seg_totals: Dict[str, int] = {}
        area_totals: Dict[tuple[str, str], int] = {}
        unit_totals: Dict[tuple[str, str, str], int] = {}

        for r in records:
            if r.headcount <= 0:
                continue
            seg_totals[r.segmento_des] = seg_totals.get(r.segmento_des, 0) + r.headcount
            area_totals[(r.segmento_des, r.area_des)] = (
                area_totals.get((r.segmento_des, r.area_des), 0) + r.headcount
            )
            unit_totals[(r.segmento_des, r.area_des, r.unidade_des)] = r.headcount

        base_total_company = total_company
        base_seg_totals = dict(seg_totals)
        base_area_totals = dict(area_totals)

        # Descoberta e Reconciliação 100% Automática em Memória (PPTX <-> Base Excel/Databricks)
        self.config = StructureSyncService.reconcile(prs, records, self.base_config)
        sync_notes: List[str] = list(self.config.pop("_sync_notes", []))

        protected_paths: Dict[int, set] = {}
        for _k in ("date_shapes", "company_total_shapes", "segment_total_shapes", "area_subtotal_shapes"):
            for _e in self.config.get(_k, []):
                protected_paths.setdefault(int(_e["slide_index"]), set()).add(str(_e["shape_path"]).strip())
        for _u in self.config.get("unit_boxes", []):
            if not _u.get("force_remove"):
                _s = protected_paths.setdefault(int(_u["slide_index"]), set())
                _s.add(str(_u["name_path"]).strip())
                _s.add(str(_u["badge_path"]).strip())

        active_area_keys = {(r.segmento_des, r.area_des) for r in records if r.headcount > 0}

        unit_lookup: Dict[tuple[int, str], Dict[str, Any]] = {}
        active_databricks_keys = set()

        for u in self.config.get("unit_boxes", []):
            s_i = int(u["slide_index"])
            full_k = f"{u['name_path']} | {u['badge_path']}"
            unit_lookup[(s_i, full_k)] = u
            unit_lookup[(s_i, str(u["badge_path"]).strip())] = u
            unit_lookup[(s_i, str(u["name_path"]).strip())] = u

            u_tuple = (u["segmento"], u["area"], u["unidade"])
            hc_val = unit_totals.get(u_tuple, 0)
            if hc_val > 0 and not u.get("force_remove"):
                active_databricks_keys.add((s_i, full_k))

        area_sub_lookup: Dict[tuple[int, str], Dict[str, Any]] = {}
        for a in self.config.get("area_subtotal_shapes", []):
            area_sub_lookup[(int(a["slide_index"]), str(a["shape_path"]).strip())] = a

        seg_tot_lookup: Dict[tuple[int, str], Dict[str, Any]] = {}
        for s in self.config.get("segment_total_shapes", []):
            seg_tot_lookup[(int(s["slide_index"]), str(s["shape_path"]).strip())] = s

        unit_overrides_by_key: Dict[tuple[int, str], Dict[str, Any]] = {}
        removed_unit_keys: Dict[tuple[int, str], Dict[str, Any]] = {}
        direct_shape_overrides: Dict[tuple[int, str], Dict[str, Any]] = {}
        unmapped_overrides: List[Dict[str, Any]] = []
        custom_added_boxes: List[Dict[str, Any]] = []
        custom_removed_unmapped: List[Dict[str, Any]] = []

        cascaded_area_deltas: Dict[tuple[str, str], int] = {}
        cascaded_seg_deltas: Dict[str, int] = {}
        cascaded_company_delta: int = 0

        def _apply_cascade_delta(seg_name: Optional[str], area_name: Optional[str], delta_val: int) -> None:
            nonlocal total_company, cascaded_company_delta
            if delta_val == 0:
                return
            if seg_name and area_name:
                a_tuple = (seg_name, area_name)
                area_totals[a_tuple] = max(0, area_totals.get(a_tuple, 0) + delta_val)
                cascaded_area_deltas[a_tuple] = cascaded_area_deltas.get(a_tuple, 0) + delta_val
            if seg_name:
                seg_totals[seg_name] = max(0, seg_totals.get(seg_name, 0) + delta_val)
                cascaded_seg_deltas[seg_name] = cascaded_seg_deltas.get(seg_name, 0) + delta_val
            total_company = max(0, total_company + delta_val)
            cascaded_company_delta += delta_val

        for u in self.config.get("unit_boxes", []):
            s_i = int(u["slide_index"])
            full_k = f"{u['name_path']} | {u['badge_path']}"
            u_tuple = (u["segmento"], u["area"], u["unidade"])
            hc_val = unit_totals.get(u_tuple, 0)
            if hc_val == 0 or u.get("force_remove") or (s_i, full_k) not in active_databricks_keys:
                removed_unit_keys[(s_i, full_k)] = {"auto_purge": True, "unidade": u["unidade"]}

        if custom_overrides:
            for ov in custom_overrides:
                try:
                    s_idx = int(ov["slide_index"]) - 1
                except (KeyError, ValueError, TypeError):
                    continue
                if not (0 <= s_idx < len(prs.slides)):
                    continue

                action = str(ov.get("action") or "UPDATE").strip().upper()
                shape_key = str(ov.get("shape_id", "")).strip()
                raw_num = ov.get("new_number")
                has_num = raw_num is not None and str(raw_num).strip() != ""
                new_num_int: Optional[int] = None
                if has_num:
                    try:
                        new_num_int = int(round(float(str(raw_num).strip())))
                    except ValueError:
                        new_num_int = None

                if action == "ADD":
                    add_hc = new_num_int if new_num_int is not None else 0
                    u_ref = unit_lookup.get((s_idx, shape_key))
                    if u_ref is not None:
                        inf_seg, inf_area = u_ref["segmento"], u_ref["area"]
                    else:
                        inf_seg, inf_area = self._infer_segment_and_area_by_column(
                            prs.slides[s_idx], s_idx, shape_key
                        )
                    _apply_cascade_delta(inf_seg, inf_area, add_hc)
                    custom_added_boxes.append(
                        {
                            "slide_index": s_idx,
                            "template_shape_key": shape_key,
                            "new_area_name": (ov.get("new_area_name") or "Nova Unidade").strip(),
                            "new_manager": (ov.get("new_manager") or "(Definir Gestor)").strip(),
                            "new_headcount": add_hc,
                            "segmento": inf_seg or "",
                            "area": inf_area or "",
                        }
                    )
                    continue

                if action == "REMOVE":
                    if (s_idx, shape_key) in unit_lookup:
                        u_rule = unit_lookup[(s_idx, shape_key)]
                        canonical_key = f"{u_rule['name_path']} | {u_rule['badge_path']}"
                        removed_unit_keys[(s_idx, canonical_key)] = ov
                        u_tuple = (u_rule["segmento"], u_rule["area"], u_rule["unidade"])
                        prev_hc = unit_totals.get(u_tuple, 0)
                        unit_totals[u_tuple] = 0
                        _apply_cascade_delta(u_rule["segmento"], u_rule["area"], -prev_hc)
                    else:
                        custom_removed_unmapped.append(ov)
                    continue

                if (s_idx, shape_key) in unit_lookup:
                    u_rule = unit_lookup[(s_idx, shape_key)]
                    canonical_key = f"{u_rule['name_path']} | {u_rule['badge_path']}"
                    unit_overrides_by_key[(s_idx, canonical_key)] = ov
                    if new_num_int is not None:
                        u_tuple = (u_rule["segmento"], u_rule["area"], u_rule["unidade"])
                        prev_hc = unit_totals.get(u_tuple, 0)
                        delta = new_num_int - prev_hc
                        unit_totals[u_tuple] = new_num_int
                        _apply_cascade_delta(u_rule["segmento"], u_rule["area"], delta)
                    continue

                if (s_idx, shape_key) in area_sub_lookup:
                    a_rule = area_sub_lookup[(s_idx, shape_key)]
                    direct_shape_overrides[(s_idx, shape_key)] = ov
                    if new_num_int is not None:
                        a_tuple = (a_rule["segmento"], a_rule["area"])
                        prev_sub = area_totals.get(a_tuple, 0)
                        delta = new_num_int - prev_sub
                        _apply_cascade_delta(a_rule["segmento"], a_rule["area"], delta)
                    continue

                if (s_idx, shape_key) in seg_tot_lookup:
                    s_rule = seg_tot_lookup[(s_idx, shape_key)]
                    direct_shape_overrides[(s_idx, shape_key)] = ov
                    if new_num_int is not None:
                        seg_name = s_rule["segmento"]
                        prev_seg = seg_totals.get(seg_name, 0)
                        delta = new_num_int - prev_seg
                        _apply_cascade_delta(seg_name, None, delta)
                    continue

                unmapped_overrides.append(ov)

        report = ExecutionReport(
            original_filename=original_filename,
            output_filename=f"Atualizado_{original_filename}",
            slides_before=slides_before,
            slides_after=slides_before,
            total_company_headcount=total_company,
            reference_label=reference_label,
        )

        manual_entries: List[ChangeLogEntry] = []
        cascade_entries: List[ChangeLogEntry] = []
        regular_entries: List[ChangeLogEntry] = []
        date_entries: List[ChangeLogEntry] = []

        # 1. Datas Base
        for d in self.config.get("date_shapes", []):
            sh = PptxSurgicalEngine.find_shape_by_path(prs.slides[d["slide_index"]], d["shape_path"])
            if sh:
                old_v, new_v = PptxSurgicalEngine.update_numeric_badge(sh, reference_label)
                st = UpdateStatus.UPDATED_NUMBER if old_v != new_v else UpdateStatus.UNCHANGED
                date_entries.append(
                    ChangeLogEntry(
                        d["slide_index"] + 1,
                        str(d["shape_path"]),
                        "Tag de Competência",
                        TargetType.REFERENCE_DATE,
                        old_v,
                        new_v,
                        st,
                        "Tag de competência",
                    )
                )

        if len(prs.slides) > 6:
            obs_sh = PptxSurgicalEngine.find_shape_by_path(prs.slides[6], "38")
            if obs_sh and getattr(obs_sh, "has_text_frame", False):
                old_obs = obs_sh.text_frame.text.strip()
                if "quadro de maio" in old_obs.lower():
                    PptxSurgicalEngine.update_numeric_badge(obs_sh, "")

        # 2. Total Geral CNI
        for c in self.config.get("company_total_shapes", []):
            sh = PptxSurgicalEngine.find_shape_by_path(prs.slides[c["slide_index"]], c["shape_path"])
            if sh:
                old_v, new_v = PptxSurgicalEngine.update_numeric_badge(sh, total_company)
                if cascaded_company_delta != 0:
                    sign = "+" if cascaded_company_delta > 0 else ""
                    cascade_entries.append(
                        ChangeLogEntry(
                            c["slide_index"] + 1,
                            str(c["shape_path"]),
                            "Total Geral CNI",
                            TargetType.HEADCOUNT_COMPANY_TOTAL,
                            f"Total [{old_v}]",
                            f"Total [{new_v}]",
                            UpdateStatus.UPDATED_NUMBER,
                            f"🔄 Recalculado em cascata: Base ({base_total_company}) {sign}{cascaded_company_delta} = {total_company}",
                        )
                    )
                else:
                    st = UpdateStatus.UPDATED_NUMBER if old_v != new_v else UpdateStatus.UNCHANGED
                    regular_entries.append(
                        ChangeLogEntry(
                            c["slide_index"] + 1,
                            str(c["shape_path"]),
                            "Total Geral CNI",
                            TargetType.HEADCOUNT_COMPANY_TOTAL,
                            f"Total [{old_v}]",
                            f"Total [{new_v}]",
                            st,
                            "Total Geral da Confederação",
                        )
                    )

        # 3. Totais de Diretoria / Entidade
        for s in self.config.get("segment_total_shapes", []):
            s_idx = int(s["slide_index"])
            s_path = str(s["shape_path"]).strip()
            seg_name = s["segmento"]
            sh = PptxSurgicalEngine.find_shape_by_path(prs.slides[s_idx], s_path)
            if sh and seg_name in seg_totals:
                old_v, new_v = PptxSurgicalEngine.update_numeric_badge(sh, seg_totals[seg_name])
                st = UpdateStatus.UPDATED_NUMBER if old_v != new_v else UpdateStatus.UNCHANGED
                regular_entries.append(
                    ChangeLogEntry(
                        s_idx + 1,
                        s_path,
                        f"Total Diretoria · {seg_name}",
                        TargetType.HEADCOUNT_SEGMENT_TOTAL,
                        f"Total [{old_v}]",
                        f"Total [{new_v}]",
                        st,
                        f"Total Diretoria ({seg_name})",
                    )
                )

        # 4. Subtotais de Superintendência
        for a in self.config.get("area_subtotal_shapes", []):
            s_idx = int(a["slide_index"])
            a_path = str(a["shape_path"]).strip()
            seg_name = a["segmento"]
            area_name = a["area"]
            sh = PptxSurgicalEngine.find_shape_by_path(prs.slides[s_idx], a_path)
            if sh and (area_name == "__EXTINCT_AREA__" or ((seg_name, area_name) not in active_area_keys and (seg_name, area_name) not in cascaded_area_deltas)):
                old_sub = sh.text_frame.text.strip() if getattr(sh, "has_text_frame", False) else ""
                PptxSurgicalEngine.remove_shapes_with_backdrop(prs.slides[s_idx], [a_path])
                manual_entries.append(
                    ChangeLogEntry(
                        s_idx + 1,
                        a_path,
                        f"Subtotal Superintendência · {area_name}",
                        TargetType.HEADCOUNT_AREA_SUBTOTAL,
                        f"Subtotal [{old_sub}]",
                        f"(Removido: área sem registros em {reference_label}) [0]",
                        UpdateStatus.REMOVED_BOX,
                        f"🗑️ Área extinta: selo de subtotal removido · {seg_name}",
                    )
                )
                continue
            if sh:
                val = area_totals.get((seg_name, area_name), 0)
                old_v, new_v = PptxSurgicalEngine.update_numeric_badge(sh, val)
                st = UpdateStatus.UPDATED_NUMBER if old_v != new_v else UpdateStatus.UNCHANGED
                regular_entries.append(
                    ChangeLogEntry(
                        s_idx + 1,
                        a_path,
                        f"Subtotal Superintendência · {area_name}",
                        TargetType.HEADCOUNT_AREA_SUBTOTAL,
                        f"Subtotal [{old_v}]",
                        f"Subtotal [{new_v}]",
                        st,
                        f"Subtotal Superintendência · {area_name}",
                    )
                )

        # 5. Purga Real de Unidades Extintas / Ausentes + Atualização das Unidades Ativas
        for u in list(self.config.get("unit_boxes", [])):
            s_idx = int(u["slide_index"])
            sl = prs.slides[s_idx]
            canonical_key = f"{u['name_path']} | {u['badge_path']}"

            if (s_idx, canonical_key) in removed_unit_keys:
                rem_info = PptxSurgicalEngine.remove_unit_group_in_column(
                    sl,
                    canonical_key,
                    compact_column=auto_layout,
                    protected_paths=protected_paths.get(s_idx),
                )
                old_txt = rem_info["old_text"] if rem_info and rem_info.get("old_text") else u["unidade"]
                old_hc_val = rem_info["old_hc"] if rem_info else 0
                manual_entries.append(
                    ChangeLogEntry(
                        s_idx + 1,
                        canonical_key,
                        u["unidade"],
                        TargetType.HEADCOUNT_UNIT,
                        f"{old_txt} [{old_hc_val}]",
                        f"(Excluída Automaticamente: ausente em {reference_label}) [0]",
                        UpdateStatus.REMOVED_BOX,
                        f"🗑️ Unidade extinta/ausente na base ({reference_label}) excluída do slide e coluna alinhada · {u['segmento']}",
                    )
                )
                continue

            name_sh = PptxSurgicalEngine.find_shape_by_path(sl, u["name_path"])
            badge_sh = PptxSurgicalEngine.find_shape_by_path(sl, u["badge_path"])
            if not name_sh or not badge_sh:
                continue

            ov = unit_overrides_by_key.get((s_idx, canonical_key))
            val = unit_totals.get((u["segmento"], u["area"], u["unidade"]), 0)
            old_num, new_num = PptxSurgicalEngine.update_numeric_badge(badge_sh, val)
            old_name = name_sh.text_frame.text.strip().replace("\x0b", " ").replace("\n", " ")
            new_name = old_name

            if ov is not None:
                if ov.get("new_area_name") or ov.get("new_manager"):
                    curr_area, _ = PptxSurgicalEngine.split_area_and_manager(name_sh)
                    target_area = ov.get("new_area_name") or u.get("rename_to") or curr_area
                    _, new_name = PptxSurgicalEngine.update_area_name_preserving_manager(
                        name_sh, target_area, new_manager_override=ov.get("new_manager")
                    )
                manual_entries.append(
                    ChangeLogEntry(
                        s_idx + 1,
                        canonical_key,
                        u["unidade"],
                        TargetType.HEADCOUNT_UNIT,
                        f"{old_name} [{old_num}]",
                        f"{new_name} [{new_num}]",
                        UpdateStatus.UPDATED_NAME_AND_NUMBER,
                        f"✏️ Editado manualmente · {u['segmento']} › {u['area']}",
                    )
                )
            else:
                if u.get("rename_to"):
                    old_name, new_name = PptxSurgicalEngine.update_area_name_preserving_manager(
                        name_sh, u["rename_to"]
                    )
                    u["is_new"] = True
                    PptxSurgicalEngine.apply_dashed_border_xml(name_sh, hex_color="164194")
                    st = UpdateStatus.UPDATED_NAME_AND_NUMBER
                else:
                    if "instituto integrador" in normalize_text_key(old_name):
                        u["is_new"] = True
                        PptxSurgicalEngine.apply_dashed_border_xml(name_sh, hex_color="164194")
                    st = UpdateStatus.UPDATED_NUMBER if old_num != new_num else UpdateStatus.UNCHANGED

                regular_entries.append(
                    ChangeLogEntry(
                        s_idx + 1,
                        canonical_key,
                        u["unidade"],
                        TargetType.HEADCOUNT_UNIT,
                        f"{old_name} [{old_num}]",
                        f"{new_name} [{new_num}]",
                        st,
                        f"Unidade sincronizada · {u['segmento']} › {u['area']}",
                    )
                )

        # 6. Novas Caixas Clonadas Dinamicamente
        for cl in self.config.get("cloned_new_boxes", []):
            s_idx = int(cl["slide_index"])
            sl = prs.slides[s_idx]
            val = unit_totals.get((cl["segmento"], cl["area"], cl["unidade"]), 0)
            if val <= 0:
                continue

            role_code = cl.get("role_code") or classify_organizational_node(cl["segmento"], cl["area"], cl["unidade"])[0]
            cloned_info = PptxSurgicalEngine.clone_unit_group_in_column(
                slide=sl,
                template_group_id=int(cl["template_group_id"]),
                new_area_name=cl["display_name"],
                new_headcount=val,
                manager_placeholder=cl.get("manager_placeholder", "(Definir Gestor)"),
                left_override=cl.get("left_override"),
                top_override=cl.get("top_override"),
                role_code=role_code,
            )

            self.config.setdefault("unit_boxes", []).append(
                {
                    "slide_index": s_idx,
                    "name_path": cloned_info["name_path"],
                    "badge_path": cloned_info["badge_path"],
                    "segmento": cl["segmento"],
                    "area": cl["area"],
                    "unidade": cl["unidade"],
                    "role_code": role_code,
                    "is_cloned": True,
                    "is_new": True,
                    "force_remove": False,
                }
            )

            regular_entries.append(
                ChangeLogEntry(
                    s_idx + 1,
                    f"{cloned_info['name_path']} | {cloned_info['badge_path']}",
                    cl["display_name"],
                    TargetType.HEADCOUNT_UNIT,
                    "(Não existia no slide)",
                    f"{cl['display_name']} {cl.get('manager_placeholder', '(Definir Gestor)')} [{val}]",
                    UpdateStatus.CLONED_NEW_BOX,
                    f"➕ Nova estrutura criada automaticamente (borda pontilhada) · {cl['segmento']} › {cl['area']}",
                )
            )

        # 6B. Novas Caixas Adicionadas Manualmente na Web (action == 'ADD')
        for add_req in custom_added_boxes:
            s_idx = int(add_req["slide_index"])
            sl = prs.slides[s_idx]
            try:
                role_code, _ = classify_organizational_node(
                    add_req["segmento"], add_req["area"], add_req["new_area_name"]
                )
                cloned_info = PptxSurgicalEngine.clone_unit_group_in_column(
                    slide=sl,
                    template_group_id=str(add_req["template_shape_key"]),
                    new_area_name=add_req["new_area_name"],
                    new_headcount=add_req["new_headcount"],
                    manager_placeholder=add_req["new_manager"],
                    role_code=role_code,
                )
                self.config.setdefault("unit_boxes", []).append(
                    {
                        "slide_index": s_idx,
                        "name_path": cloned_info["name_path"],
                        "badge_path": cloned_info["badge_path"],
                        "segmento": add_req["segmento"],
                        "area": add_req["area"],
                        "unidade": add_req["new_area_name"],
                        "role_code": role_code,
                        "is_cloned": True,
                        "is_new": True,
                        "force_remove": False,
                    }
                )
                manual_entries.append(
                    ChangeLogEntry(
                        s_idx + 1,
                        f"{cloned_info['name_path']} | {cloned_info['badge_path']}",
                        add_req["new_area_name"],
                        TargetType.HEADCOUNT_UNIT,
                        "(Adicionada via Web)",
                        f"{add_req['new_area_name']} {add_req['new_manager']} [{add_req['new_headcount']}]",
                        UpdateStatus.CLONED_NEW_BOX,
                        f"➕ Nova caixa criada pelo utilizador · {add_req['segmento']} › {add_req['area']}",
                    )
                )
            except Exception:
                pass

        report.changes = manual_entries + cascade_entries + regular_entries + date_entries

        output_buffer = io.BytesIO()
        prs.save(output_buffer)
        pptx_bytes = output_buffer.getvalue()

        prs_val = Presentation(io.BytesIO(pptx_bytes))
        report.slides_after = len(prs_val.slides)
        report.technical_validation_passed = report.slides_before == report.slides_after
        report.updated_count = sum(1 for c in report.changes if c.status != UpdateStatus.UNCHANGED)
        report.unchanged_count = sum(1 for c in report.changes if c.status == UpdateStatus.UNCHANGED)
        report.conflict_count = sum(1 for c in report.changes if c.status == UpdateStatus.REMOVED_BOX)
        report.validation_notes.extend(sync_notes)
        report.slides_preview = PptxSurgicalEngine.extract_slides_preview(prs_val, self.config)

        return pptx_bytes, report